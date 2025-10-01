# These imports will still use the correct versions (python-docx and fpdf2)
import importlib.util
import openai
# from openai import OpenAI
from docx import Document
from fpdf import FPDF
import os
# from dotenv import load_dotenv
# Optional import below
import importlib

try:
    import streamlit as st
except ImportError:
    st = None

# Try to load .env only if available
dotenv_spec = importlib.util.find_spec("dotenv")
if dotenv_spec is not None:
    from dotenv import load_dotenv
    load_dotenv()


# Handle Streamlit Secrets gracefully
def get_api_key():
    if st and "OPENAI_API_KEY" in st.secrets:
        return st.secrets["OPENAI_API_KEY"]

    api_key = os.getenv("OPENAI_API_KEY")
    if api_key:
        return api_key

    raise ValueError("OPENAI_API_KEY not set. Add it to .env, GitHub secrets, or Streamlit secrets.")


OPENAI_API_KEY = get_api_key()
openai.api_key = OPENAI_API_KEY

# VOID
# Insert OpenAI API key here when actually deploying
# client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))


def generate_policy(company_name: str, industry: str, jurisdiction: str, size: str, policy_type: str) -> str:
    prompt = f"""
    You are an experienced compliance consultant specializing in {jurisdiction} law.

    Generate a formal {policy_type} tailored for the following company:
    - Company name: {company_name}
    - Industry: {industry}
    - Jurisdiction: {jurisdiction}
    - Company size: {size}

    The policy must follow this structure:

    1. Title
    2. Purpose
    3. Scope
    4. Responsibilities
    5. Procedures
    6. Review cycle

    Requirements:
    - Use clear, professional, and legally appropriate language.
    - Adapt complexity to the company size.
    - Ensure the content aligns with {jurisdiction}-specific compliance expectations.
    - Write in polished prose suitable for a corporate handbook.
    """

    response = openai.ChatCompletion.create(
        model="gpt-4o-mini",
        messages=[
            {"role": "system", "content": "You are a compliance consultant generating policies for businesses."},
            {"role": "user", "content": prompt}
        ],
        temperature=0.5
    )
    return response.choices[0].message.content


def save_to_docx(policy_text: str, filename: str) -> None:
    doc = Document()
    doc.add_heading("Generated Compliance Policy", level=0)
    doc.add_paragraph(policy_text)
    doc.add_paragraph(
        "\nDisclaimer: This document is AI-Generated and does not constitute legal advice."
    )
    doc.save(filename)


def save_to_pdf(policy_text: str, filename: str) -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.multi_cell(0, 10, "Generated Compliance Policy", align="C")
    pdf.ln(10)

    for line in policy_text.split("\n"):
        pdf.multi_cell(0, 10, line.strip())

    pdf.ln(10)
    pdf.set_font("Arial", "I", 10)
    pdf.multi_cell(0, 10, "Disclaimer: This document is AI-Generated and does not constitute legal advice.")
    pdf.output(filename)
